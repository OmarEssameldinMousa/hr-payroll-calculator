"""Export a PeriodResult into the exact hr_report_V1 layout.

Opens the bundled template (assets/hr_report_template.xlsx), keeps the
two header rows of each sheet, removes the sample rows, and writes data
rows by cloning the sample row styles (row 3 = data, row 5 = totals).
Column K (ملحوظة) holds the raw status text from the device export.
"""
from __future__ import annotations

import io
import shutil
import subprocess
import tempfile
from copy import copy
from datetime import datetime
from fractions import Fraction
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .engine import Employee, PeriodResult

TEMPLATE = Path(__file__).resolve().parent.parent / "assets" / "hr_report_template.xlsx"
NOFILL = PatternFill(fill_type=None)
YELLOW = PatternFill("solid", fgColor="FFFFE599")  # template colour for cells HR fixed
EDITED_COL = {"clock_out": 12, "clock_in": 13, "status": 13}
STATUS_LABEL = {
    "absent_unexcused": "absent",
    "absent_excused": "absent (excused)",
    "off_day": "off day",
}


# ---------------------------------------------------------------- formatting
def fmt_days(x: Fraction, zero: str = "_") -> str:
    if not x:
        return zero
    sign = "-" if x < 0 else ""
    a = abs(x)
    whole, rem = a.numerator // a.denominator, a - a.numerator // a.denominator
    body = f"{whole},{rem}" if whole and rem else (f"{whole}" if whole else f"{rem}")
    return f"{sign}{body} {'day' if a <= 1 else 'days'}"


def fmt_hm(minutes: int | None, zero: str = "_") -> str:
    if not minutes:
        return zero
    sign = "-" if minutes < 0 else ""
    m = abs(int(minutes))
    return f"{sign}{m // 60:02d}:{m % 60:02d} hrs"


def fmt_clock(dt: datetime | None, work_date) -> str:
    if dt is None:
        return "_"
    s = dt.strftime("%I:%M") + dt.strftime("%p").lower() + ","
    if dt.date() != work_date:
        s += f" {dt.day}/{dt.strftime('%b').lower()}"
    return s


# ---------------------------------------------------------------- writer
def _clear_below(ws, first_row: int):
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= first_row:
            ws.unmerge_cells(str(rng))
    if ws.max_row >= first_row:
        ws.delete_rows(first_row, ws.max_row - first_row + 1)


def build_report(res: PeriodResult, emp: Employee, include_summary: bool = False) -> bytes:
    wb = load_workbook(TEMPLATE)
    ws = wb["attendance"]
    ncol = 15  # A..O
    data_style = [copy(ws.cell(3, c)._style) for c in range(1, ncol + 1)]
    total_style = [copy(ws.cell(5, c)._style) for c in range(1, ncol + 1)]
    date_fmt = ws.cell(3, 14).number_format
    ws.cell(2, 11).value = "ملحوظة"
    _clear_below(ws, 3)

    row = 3
    for r in res.days:
        d = r.day
        label = STATUS_LABEL.get(d.status)
        cin = label if label else fmt_clock(d.clock_in, d.date)
        cout = label if label else fmt_clock(d.clock_out, d.date)
        if d.status == "additional_day":
            cin = f"{cin} (additional)"
        vals = {
            1: fmt_days(r.total_deduct),
            2: "_",
            3: "_",
            4: f'"+{r.exception_earned}min"' if r.exception_earned else "_",
            5: fmt_hm(r.overtime_minutes),
            6: fmt_days(r.lateness_deduct),
            7: fmt_hm(r.late_minutes),
            8: "_",
            9: "_",
            10: fmt_hm(r.worked_minutes),
            11: d.raw_status or "_",
            12: cout,
            13: cin,
            14: datetime.combine(d.date, datetime.min.time()),
            15: d.day_name,
        }
        for c in range(1, ncol + 1):
            cell = ws.cell(row, c, vals[c])
            cell._style = copy(data_style[c - 1])
            cell.fill = NOFILL
        ws.cell(row, 14).number_format = date_fmt
        for f in d.edited:
            ws.cell(row, EDITED_COL[f]).fill = YELLOW
        row += 1

    k = res.kpis
    totals = {1: fmt_days(k["deductions_total"], "0 day"),
              2: "0 day",
              5: fmt_hm(k["overtime_minutes"], "00:00 hrs"),
              6: fmt_days(k["deduct_lateness"], "0 day"),
              8: "0 day",
              10: fmt_hm(k["worked_total_minutes"], "00:00 hrs")}
    for c in range(1, ncol + 1):
        cell = ws.cell(row, c, totals.get(c))
        cell._style = copy(total_style[c - 1])

    # print setup (used by Excel printing and the PDF): landscape, one page wide,
    # header rows repeated, employee + period in the page header
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_title_rows = "1:2"
    period = f"{res.days[0].day.date} → {res.days[-1].day.date}" if res.days else ""
    ws.oddHeader.center.text = f"{emp.name} ({emp.code})   {period}".replace("&", "&&")  # & is a header code

    # actions sheet is kept for layout parity; no rule produces manual actions
    _clear_below(wb["actions"], 3)

    if include_summary:
        _summary_sheet(wb, res, emp)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(res: PeriodResult, emp: Employee, include_summary: bool = False) -> bytes:
    """Build the Excel report, then convert it to PDF with LibreOffice (headless)."""
    wb = load_workbook(io.BytesIO(build_report(res, emp, include_summary)))
    del wb["actions"]   # always empty → would only add a blank page
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "report.xlsx"
        wb.save(src)
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            raise RuntimeError("LibreOffice is not installed (needed for PDF export)")
        # private profile dir so it works even while LibreOffice is open on the desktop
        proc = subprocess.run([soffice, f"-env:UserInstallation=file://{tmp}/profile", "--headless",
                               "--convert-to", "pdf", "--outdir", tmp, str(src)],
                              capture_output=True, text=True, timeout=120)
        pdf = src.with_suffix(".pdf")
        if not pdf.exists():
            raise RuntimeError(f"PDF conversion failed: {proc.stderr or proc.stdout}")
        return pdf.read_bytes()


def summary_rows(res: PeriodResult) -> list[tuple[str, str, str]]:
    k = res.kpis
    return [
        ("Time & attendance", "اجمالي عدد ساعات العمل الرسمية",
         f"{fmt_hm(k['worked_regular_minutes'], '00:00 hrs')} from {fmt_hm(k['official_minutes'], '00:00 hrs')}"),
        ("Time & attendance", "اجمالي الاوفر تايم", fmt_hm(k["overtime_minutes"], "00:00 hrs")),
        ("Time & attendance", "اجمالي رصيد ساعات العمل", fmt_hm(k["balance_minutes"], "00:00 hrs")),
        ("Attendance", "اجمالي عدد ايام الحضور", str(k["attendance_days"])),
        ("Attendance", "عدد حالات خصم ربع يوم (تأخير أكثر من 10 دقائق)", str(k["late_quarter_count"])),
        ("Attendance", "عدد حالات خصم ثلث يوم (تأخير أكثر من 30 دقيقة)", str(k["late_third_count"])),
        ("Attendance", "عدد حالات تأخير يوم الجمعة (أكثر من 30 دقيقة)", str(k["friday_late_count"])),
        ("Attendance", "عدد الاستثناءات المكتسبة من الإضافي / المستخدمة",
         f"{k['morning_exceptions_earned']} / {k['morning_exceptions_used']}"),
        ("Attendance", "عدد مرات تجاوز الاستثناء (أكثر من 30 دقيقة)", str(k["exception_exceeded_count"])),
        ("Attendance", "عدد أيام الغياب بدون إذن", f"{k['absence_unexcused']} → {fmt_days(k['deduct_absence'], '0 day')}"),
        ("Attendance", "عدد أيام الغياب بإذن", str(k["absence_excused"])),
        ("Deductible", "إجمالي الخصومات", fmt_days(k["deductions_total"], "0 day")),
    ]


def _summary_sheet(wb, res, emp):
    ws = wb.create_sheet("summary")
    hdr = Font(name="Arial", size=11, bold=True)
    grey = PatternFill("solid", fgColor="FFB7B7B7")
    ws.append(["Section", "KPI", "Value"])
    for c in ws[1]:
        c.font, c.fill = hdr, grey
    for row in summary_rows(res):
        ws.append(list(row))
    for col, w in zip("ABC", (20, 48, 28)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(horizontal="center", vertical="center")
