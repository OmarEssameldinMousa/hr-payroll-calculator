import io
from datetime import date, datetime, time, timedelta
from fractions import Fraction as Fr
from pathlib import Path

from openpyxl import load_workbook

from hr_engine.config import PAYROLL as P
from hr_engine.engine import Day, Employee, run_period
from hr_engine.parser import parse_attendance_xls
from hr_engine.report import build_report, fmt_days

D0 = date(2026, 8, 22)  # saturday; D0+5 = thu (official off), D0+6 = fri
EMP = Employee("1", "t", official_in=time(9, 30), official_out=time(17, 0), official_off_weekday=3)


def day(i, hin=None, hout=None, **kw):
    d = D0 + timedelta(days=i)
    cin = datetime.combine(d, hin) if hin else None
    cout = datetime.combine(d, hout) if hout else None
    if cin and cout and cout <= cin:
        cout += timedelta(days=1)
    return Day(d, clock_in=cin, clock_out=cout, **kw)


def test_lateness_tiers():
    r = run_period([day(0, time(9, 40), time(17)), day(1, time(9, 41), time(17)),
                    day(2, time(10, 0), time(17)), day(3, time(10, 1), time(17))], EMP, P)
    assert [x.lateness_deduct for x in r.days] == [0, Fr(-1, 4), Fr(-1, 4), Fr(-1, 3)]


def test_friday_only_over_30():
    r = run_period([day(6, time(10, 0), time(17)), day(13, time(10, 1), time(17))], EMP, P)
    assert [x.lateness_deduct for x in r.days] == [0, Fr(-1, 3)]


def test_overtime_exception_is_flat_30_minutes():
    # left 18:31 = 91m after official clock-out (> 1:00) → 30m exception next day
    days = [day(0, time(9, 30), time(18, 31)), day(1, time(10, 0), time(17)),
            day(2, time(9, 30), time(18, 0)), day(3, time(9, 45), time(17)),   # exactly 1:00 → no exception
            day(6, time(9, 30), time(1, 30)), day(7, time(10, 5), time(17))]    # 8:30 overtime → still 30m
    r = run_period(days, EMP, P)
    assert [x.exception_earned for x in r.days] == [30, 0, 0, 0, 30, 0]
    assert r.days[1].exception_used and r.days[1].lateness_deduct == 0
    assert r.days[3].lateness_deduct == Fr(-1, 4)
    assert r.days[5].lateness_deduct == Fr(-1, 3) and r.ledger[-1].rule == "EXCEPTION_EXCEEDED"


def test_absence():
    days = [Day(D0, "absent_unexcused"), Day(D0 + timedelta(1), "absent_excused"),
            Day(D0 + timedelta(2), "off_day"), Day(D0 + timedelta(5), "absent_unexcused")]  # thu = official off
    r = run_period(days, EMP, P)
    assert [x.total_deduct for x in r.days] == [-2, 0, 0, 0]


def test_overtime_is_time_after_official_clock_out():
    # 30m late but stayed 1:00 after 17:00 → overtime 60m, not 30m; exactly 1:00 earns nothing
    r = run_period([day(0, time(10, 0), time(18, 0)), day(1, time(9, 30), time(18, 1))], EMP, P)
    assert [x.overtime_minutes for x in r.days] == [60, 61]
    assert [x.exception_earned for x in r.days] == [0, 30]


def test_day_off_cancels_exception():
    days = [day(0, time(9, 30), time(19)), Day(D0 + timedelta(1), "off_day"), day(2, time(9, 50), time(17)),
            day(4, time(9, 30), time(19)), day(5, time(9, 30), time(19)),    # thu = official off
            day(6, time(9, 30), time(17)), day(7, time(9, 50), time(17))]
    r = run_period(days, EMP, P)
    assert r.days[2].exception_available == 0 and r.days[2].lateness_deduct == Fr(-1, 4)
    assert r.days[5].exception_available == 0


def test_official_day_off_is_vacation():
    r = run_period([day(5, time(11, 0), time(19, 0)), day(6, time(10, 0), time(17))], EMP, P)
    off = r.days[0]
    assert (off.total_deduct, off.overtime_minutes, off.exception_earned, off.official_minutes) == (0, 0, 0, 0)


def test_hours_balance():
    r = run_period([day(0, time(9, 30), time(17)), day(1, time(9, 30), time(19))], EMP, P)
    k = r.kpis
    assert (k["official_minutes"], k["overtime_minutes"], k["balance_minutes"]) == (900, 120, 120)


def test_fmt_days():
    assert fmt_days(Fr(-5, 4)) == "-1,1/4 days" and fmt_days(Fr(-1, 4)) == "-1/4 day" and fmt_days(0) == "_"


def test_report_note_column_has_raw_status():
    days = [Day(D0, "absent_unexcused", raw_status="غياب"), day(1, time(9, 30), time(17))]
    ws = load_workbook(io.BytesIO(build_report(run_period(days, EMP, P), EMP)))["attendance"]
    assert ws.cell(2, 11).value == "ملحوظة"
    assert (ws.cell(3, 11).value, ws.cell(4, 11).value) == ("غياب", "_")


def test_sample_file_end_to_end():
    p = parse_attendance_xls((Path(__file__).parent / "sample_attendance.xls").read_bytes())
    assert (p.period_start, p.period_end, len(p.days)) == (date(2026, 8, 22), date(2026, 9, 21), 31)
    d = {x.date: x for x in p.days}
    assert d[date(2026, 8, 25)].clock_out == datetime(2026, 8, 26, 7, 0)   # overnight
    assert d[date(2026, 9, 3)].status == "off_day"
    days = [Day(x.date, x.status, x.clock_in, x.clock_out, x.raw_status_text) for x in p.days]
    res = run_period(days, EMP, P)
    assert res.kpis["absence_unexcused"] == sum(
        1 for x in p.days if x.status == "absent_unexcused" and x.date.weekday() != 3)
    assert build_report(res, EMP)[:2] == b"PK"


def test_additional_day_does_not_earn_exception():
    days = [day(0, time(9, 30), time(19), status="additional_day"), day(1, time(9, 50), time(17))]
    r = run_period(days, EMP, P)
    assert r.days[0].exception_earned == 0 and r.days[1].lateness_deduct == Fr(-1, 4)


def test_report_highlights_cells_hr_fixed():
    fixed = day(0, time(9, 30), time(17))
    fixed.edited.add("clock_out")
    ws = load_workbook(io.BytesIO(build_report(run_period([fixed], EMP, P), EMP)))["attendance"]
    assert ws.cell(3, 12).fill.fgColor.rgb == "FFFFE599"   # clock-out: fixed
    assert ws.cell(3, 13).fill.fgColor.rgb != "FFFFE599"   # clock-in: untouched
